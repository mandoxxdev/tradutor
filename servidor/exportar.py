"""
Exportação da ata da reunião em .docx e .txt bilíngue.

Cada fala gera: data, hora, idioma falado, texto original e tradução(ões).
Os arquivos vão para a pasta atas/ com data e hora no nome.
"""

import logging
from datetime import datetime
from pathlib import Path

from . import idiomas as tab

log = logging.getLogger("gmp.exportar")


def _rotulo(codigo: str) -> str:
    i = tab.obter(codigo)
    return i.rotulo if i else codigo.upper()


def _traducoes_texto(fala: dict, separador: str = "\n") -> str:
    partes = []
    trads = fala.get("traducoes") or {}
    for codigo, texto in trads.items():
        if len(trads) == 1:
            partes.append(texto or "(sem tradução)")
        else:
            partes.append(f"[{_rotulo(codigo)}] {texto or '(sem tradução)'}")
    if fala.get("erro"):
        partes.append(f"(aviso: {fala['erro']})")
    return separador.join(partes)


def exportar_txt(historico: list[dict], idiomas_visitante: list[str], caminho: Path, iniciada_em: datetime) -> None:
    linhas = [
        "ATA DE REUNIÃO - GMP Tradutor",
        f"Início: {iniciada_em.strftime('%d/%m/%Y %H:%M')}    Exportado em: {datetime.now().strftime('%d/%m/%Y %H:%M')}",
        "Idioma da casa: " + _rotulo(tab.PORTUGUES),
        "Idioma(s) do visitante: " + (", ".join(_rotulo(c) for c in idiomas_visitante) or "-"),
        f"Total de falas: {len(historico)}",
        "=" * 78,
        "",
    ]
    for f in historico:
        aviso = "  [idioma fora da sessão]" if f.get("fora_da_sessao") else ""
        linhas.append(f"[{f.get('data', '')} {f.get('hora', '')}] {_rotulo(f.get('idioma', ''))}{aviso}")
        linhas.append(f"ORIGINAL:  {f.get('original', '')}")
        linhas.append(f"TRADUÇÃO:  {_traducoes_texto(f, separador=chr(10) + '           ')}")
        linhas.append("")
    caminho.write_text("\n".join(linhas), encoding="utf-8")


def exportar_docx(historico: list[dict], idiomas_visitante: list[str], caminho: Path, iniciada_em: datetime) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, Cm

    doc = Document()
    # Página em paisagem com margens menores: cabe mais texto na tabela
    secao = doc.sections[0]
    secao.page_width, secao.page_height = secao.page_height, secao.page_width
    secao.orientation = 1  # paisagem
    for lado in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(secao, lado, Cm(1.5))
    estilo = doc.styles["Normal"]
    estilo.font.name = "Arial"
    estilo.font.size = Pt(10)

    logo = Path(__file__).resolve().parent.parent / "web" / "logo.png"
    if logo.exists():
        par_logo = doc.add_paragraph()
        par_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        try:
            par_logo.add_run().add_picture(str(logo), width=Cm(5))
        except Exception as e:  # logo inválido não impede a ata
            log.warning("Não foi possível inserir o logo na ata: %s", e)

    titulo = doc.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = titulo.add_run("ATA DE REUNIÃO")
    r.bold = True
    r.font.size = Pt(18)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run("GMP Industriais · GMP Tradutor").italic = True

    doc.add_paragraph(f"Início da reunião: {iniciada_em.strftime('%d/%m/%Y %H:%M')}")
    doc.add_paragraph(f"Exportado em: {datetime.now().strftime('%d/%m/%Y %H:%M')}")
    doc.add_paragraph("Idioma da casa: " + _rotulo(tab.PORTUGUES))
    doc.add_paragraph("Idioma(s) do visitante: " + (", ".join(_rotulo(c) for c in idiomas_visitante) or "-"))
    doc.add_paragraph(f"Total de falas: {len(historico)}")

    tabela = doc.add_table(rows=1, cols=4)
    tabela.style = "Table Grid"
    cab = tabela.rows[0].cells
    for celula, texto in zip(cab, ("Data / hora", "Idioma", "Texto original", "Tradução")):
        celula.text = ""
        run = celula.paragraphs[0].add_run(texto)
        run.bold = True
    for f in historico:
        linha = tabela.add_row().cells
        linha[0].text = f"{f.get('data', '')}\n{f.get('hora', '')}"
        idioma = _rotulo(f.get("idioma", ""))
        if f.get("fora_da_sessao"):
            idioma += "\n(fora da sessão)"
        linha[1].text = idioma
        linha[2].text = f.get("original", "")
        linha[3].text = _traducoes_texto(f)
    larguras = (Cm(2.6), Cm(3.6), Cm(9.8), Cm(9.8))
    for row in tabela.rows:
        for celula, largura in zip(row.cells, larguras):
            celula.width = largura

    doc.add_paragraph("")
    doc.add_paragraph("Transcrição e tradução automáticas geradas localmente pelo GMP Tradutor; "
                      "revise antes de usar como documento oficial.").runs[0].italic = True
    doc.save(str(caminho))


def exportar(historico: list[dict], idiomas_visitante: list[str], pasta_atas: Path,
             iniciada_em: datetime | None = None) -> tuple[Path, Path]:
    """Gera os dois arquivos e devolve (caminho_docx, caminho_txt)."""
    pasta_atas.mkdir(parents=True, exist_ok=True)
    iniciada_em = iniciada_em or datetime.now()
    base = f"ata_{iniciada_em.strftime('%Y-%m-%d_%H-%M')}"
    docx_path = pasta_atas / f"{base}.docx"
    txt_path = pasta_atas / f"{base}.txt"
    exportar_txt(historico, idiomas_visitante, txt_path, iniciada_em)
    exportar_docx(historico, idiomas_visitante, docx_path, iniciada_em)
    log.info("Ata exportada: %s e %s", docx_path, txt_path)
    return docx_path, txt_path
